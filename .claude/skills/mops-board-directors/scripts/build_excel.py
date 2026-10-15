#!/usr/bin/env python3
"""把董事名單 JSON 轉成固定格式的 Excel（總覽／董事名單／異動紀錄／說明與來源）。

用法：python3 build_excel.py result.json -o 董事名單_xxx_1151014.xlsx
輸入格式見 ../examples/sample_input.json 與 ../references/output-format.md。
"""
import argparse
import json
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

FONT = "微軟正黑體"
NAVY = "1F3864"
YELLOW = "FFF2CC"
GREY = "595959"
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

DIRECTOR_COLS = ["公司", "職稱", "姓名", "身分", "代表法人", "生效日",
                 "依據重訊日期", "依據重訊主旨", "持股明細核對", "備註"]
DIRECTOR_WIDTHS = [9, 13, 9, 20, 26, 11, 18, 44, 12, 46]
CHANGE_COLS = ["公司", "重訊日期", "生效日", "異動類型", "內容"]
CHANGE_WIDTHS = [9, 11, 11, 14, 64]
CHECK_VALUES = {"一致", "差異", "未核對"}


def header(ws, row, cols, widths):
    for i, (name, w) in enumerate(zip(cols, widths), 1):
        c = ws.cell(row=row, column=i, value=name)
        c.font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[row].height = 26


def body_cell(ws, row, col, value, fill=None, bold=False):
    c = ws.cell(row=row, column=col, value=value)
    c.font = Font(name=FONT, size=10, bold=bold)
    c.alignment = Alignment(vertical="center", wrap_text=True)
    c.border = BORDER
    if fill:
        c.fill = PatternFill("solid", fgColor=fill)
    return c


def table_sheet(wb, title, cols, widths, rows):
    ws = wb.create_sheet(title)
    header(ws, 1, cols, widths)
    for r, rec in enumerate(rows, 2):
        for c, key in enumerate(cols, 1):
            body_cell(ws, r, c, rec.get(key) or None)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{max(len(rows) + 1, 1)}"
    return ws


def overview(wb, data):
    ws = wb.active
    ws.title = "總覽"
    show_method = data.get("show_announce_method",
                           any(c.get("重訊公告方式") for c in data["companies"]))
    cols = ["公司", "代號", "市場別", "屆次", "任期／生效", "董事長",
            "現任董事人數", "其中獨立董事", "與持股明細不一致"]
    widths = [9, 9, 16, 14, 26, 10, 12, 12, 14]
    if show_method:
        cols.append("重訊公告方式"); widths.append(26)
    cols.append("注意事項"); widths.append(60)

    ws["A1"] = data["title"]
    ws["A1"].font = Font(name=FONT, size=14, bold=True)
    ws["A2"] = (f"資料截至：{data['as_of']}（重大訊息）｜"
                f"交叉核對：{data['holdings_month']}董監事持股餘額明細")
    ws["A2"].font = Font(name=FONT, size=9, color=GREY)
    header(ws, 4, cols, widths)

    n = len(data["companies"])
    first, last = 5, 4 + n
    for i, comp in enumerate(data["companies"]):
        r = first + i
        note = comp.get("注意事項")
        fill = YELLOW if comp.get("highlight", bool(note)) else None
        a = f"A{r}"
        values = {
            "公司": comp["公司"],
            "代號": comp.get("代號") or "—",
            "市場別": comp.get("市場別"),
            "屆次": comp.get("屆次"),
            "任期／生效": comp.get("任期"),
            # 董事名單第 1 列是表頭，MATCH 從第 2 列算起，所以 +1
            "董事長": (f'=IFERROR(INDEX(董事名單!$C:$C,MATCH(1,INDEX((董事名單!$A$2:$A$2000={a})'
                     f'*(董事名單!$B$2:$B$2000="董事長"),0),0)+1),"")'),
            "現任董事人數": f"=COUNTIF(董事名單!$A:$A,{a})",
            "其中獨立董事": f'=COUNTIFS(董事名單!$A:$A,{a},董事名單!$D:$D,"獨立董事*")',
            "與持股明細不一致": f'=COUNTIFS(董事名單!$A:$A,{a},董事名單!$I:$I,"差異")',
            "重訊公告方式": comp.get("重訊公告方式"),
            "注意事項": note,
        }
        for c, key in enumerate(cols, 1):
            cell = body_cell(ws, r, c, values[key], fill=fill)
            if key in ("代號", "現任董事人數", "其中獨立董事", "與持股明細不一致"):
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.row_dimensions[r].height = 26

    total = last + 1
    ws.cell(row=total, column=1, value="合計").font = Font(name=FONT, size=10, bold=True)
    for key in ("現任董事人數", "其中獨立董事"):
        col = cols.index(key) + 1
        L = get_column_letter(col)
        c = ws.cell(row=total, column=col, value=f"=SUM({L}{first}:{L}{last})")
        c.font = Font(name=FONT, size=10, bold=True)
        c.alignment = Alignment(horizontal="center")
    legend = ws.cell(row=total + 1, column=1,
                     value="黃底：有需要留意的事項（任期屆滿未見改選、席次出缺、重訊與持股明細不一致等）。")
    legend.font = Font(name=FONT, size=9, color=GREY)
    ws.freeze_panes = "B5"


def default_notes(data):
    subs = [c["公司"] for c in data["companies"]
            if (c.get("重訊公告方式") or "").startswith("由")]
    notes = [
        ["資料來源", "公開資訊觀測站「重大訊息」（t05st01）：https://mops.twse.com.tw/mops/#/web/t05st01，"
                 "以公司代號＋發言日期即可回查原文"],
        ["交叉核對", f"公開資訊觀測站「董事、監察人、經理人及大股東持股餘額明細」（stapap1），{data['holdings_month']}資料"],
        ["推導方法", "以最近一次全面改選（子公司為金控重新指派）之名單為基礎，依時間順序套用其後所有董事異動、"
                 "改派、辭任、解任、逝世公告，再與持股明細比對"],
    ]
    if data.get("query_period"):
        notes.append(["查詢期間", data["query_period"]])
    if subs:
        notes.append(["子公司重訊", "、".join(subs) + "本身無重訊或極少，係由母公司代為公告，已逐則查核"])
    notes += [
        ["職稱", "董事長、副董事長、常務董事依重訊；重訊未載明者依持股明細（已於備註註明）"],
        ["日期", "皆為民國紀年"],
    ]
    return notes


def validate(data):
    errs = []
    names = {c["公司"] for c in data["companies"]}
    for i, d in enumerate(data["directors"]):
        if d.get("公司") not in names:
            errs.append(f"directors[{i}] 公司「{d.get('公司')}」不在 companies 清單")
        if d.get("持股明細核對") not in CHECK_VALUES:
            errs.append(f"directors[{i}] {d.get('姓名')} 持股明細核對須為 一致／差異／未核對")
    for c in data["companies"]:
        chairs = [d for d in data["directors"] if d["公司"] == c["公司"] and d.get("職稱") == "董事長"]
        if len(chairs) != 1:
            print(f"[提醒] {c['公司']} 董事長有 {len(chairs)} 位（出缺或代理請在注意事項說明）", file=sys.stderr)
    if errs:
        sys.exit("輸入資料有誤：\n" + "\n".join(errs))


def main():
    p = argparse.ArgumentParser()
    p.add_argument("input")
    p.add_argument("-o", "--output", required=True)
    args = p.parse_args()
    data = json.load(open(args.input, encoding="utf-8"))
    validate(data)

    wb = Workbook()
    overview(wb, data)
    table_sheet(wb, "董事名單", DIRECTOR_COLS, DIRECTOR_WIDTHS, data["directors"])
    changes = sorted(data.get("changes", []),
                     key=lambda x: (order_of(data, x["公司"]), x.get("生效日") or "", x.get("重訊日期") or ""))
    table_sheet(wb, "異動紀錄", CHANGE_COLS, CHANGE_WIDTHS, changes)
    notes = data.get("notes") or default_notes(data)
    table_sheet(wb, "說明與來源", ["項目", "說明"], [14, 110],
                [{"項目": k, "說明": v} for k, v in notes])
    wb.save(args.output)
    print(f"→ {args.output}：{len(data['companies'])} 家、{len(data['directors'])} 位董事、"
          f"{len(changes)} 筆異動", file=sys.stderr)


def order_of(data, name):
    for i, c in enumerate(data["companies"]):
        if c["公司"] == name:
            return i
    return 999


if __name__ == "__main__":
    main()
